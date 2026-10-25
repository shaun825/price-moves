"""Builds public/data/statssa-history.json from Stats SA's 'CPI Average Prices' Excel files.

Usage: python3 scripts/build-statssa.py "<All urban .xlsx>" "<Provinces .xlsx>" "<master item list .xlsx>"
Only items whose product AND pack size match a Stats SA series are included (see MAP).
Stats SA terms: credit Stats SA; analysis and presentation are Price Moves' own; data not sold.
"""
import json, sys, datetime
from openpyxl import load_workbook

# our item id: (Stats SA H03 code, Stats SA size label)
MAP = {
 "GRO-001": ("01113101", "700 Gram"), "GRO-002": ("01113102", "700 Gram"), "GRO-006": ("01113104", "500 Gram"),
 "GRO-007": ("01113105", "200 Gram"), "GRO-010": ("01112601", "2,5 Kilogram"), "GRO-011": ("01112601", "5 Kilogram"),
 "GRO-015": ("01111201", "2 Kilogram"), "GRO-018": ("01112101", "2.5 Kilogram"), "GRO-019": ("01112102", "2.5 Kilogram"),
 "GRO-022": ("01112301", "1 Kilogram"), "GRO-023": ("01114002", "1 Kilogram"), "GRO-024": ("01115001", "500 Gram"),
 "GRO-025": ("01115002", "500 Gram"), "GRO-027": ("01114001", "500 Gram"), "GRO-031": ("01122404", "2 Kilogram"),
 "GRO-037": ("01122101", "Per Kilogram"), "GRO-038": ("01122106", "Per Kilogram"), "GRO-039": ("01122107", "Per Kilogram"),
 "GRO-040": ("01122108", "Per Kilogram"), "GRO-041": ("01124001", "Per Kilogram"), "GRO-042": ("01122302", "Per Kilogram"),
 "GRO-043": ("01122305", "Per Kilogram"), "GRO-044": ("01122201", "Per Kilogram"), "GRO-046": ("01125102", "Per Kilogram"),
 "GRO-047": ("01125103", "1 Kilogram"), "GRO-048": ("01125104", "500 Gram"), "GRO-049": ("01125105", "1 Kilogram"),
 "GRO-052": ("01132902", "400 Gram"), "GRO-053": ("01132901", "170 Gram"), "GRO-054": ("01125201", "300 Gram"),
 "GRO-055": ("01131401", "800 Gram"), "GRO-056": ("01133901", "400 Gram"), "GRO-057": ("01141101", "1 Litre"),
 "GRO-058": ("01141101", "2 Litre"), "GRO-059": ("01142001", "2 Litre"), "GRO-060": ("01141102", "1 Litre"),
 "GRO-062": ("01146003", "2 Kilogram"), "GRO-063": ("01146001", "1 Kilogram"), "GRO-064": ("01146002", "1 Kilogram"),
 "GRO-065": ("01152101", "500 Gram"), "GRO-066": ("01153002", "500 Gram"), "GRO-067": ("01153001", "1 Kilogram"),
 "GRO-068": ("01145001", "Per Kilogram"), "GRO-073": ("01148101", "For 18"), "GRO-075": ("01143201", "400 Gram"),
 "GRO-077": ("01163101", "1,5 Kilogram"), "GRO-078": ("01161201", "Per Kilogram"), "GRO-079": ("01162301", "Per Kilogram"),
 "GRO-081": ("01163201", "Per Kilogram"), "GRO-083": ("01161101", "Each"), "GRO-093": ("01172401", "Per Kilogram"),
 "GRO-094": ("01174101", "1 Kilogram"), "GRO-095": ("01171201", "Each"), "GRO-096": ("01171501", "Per Bunch"),
 "GRO-098": ("01172501", "Per Kilogram"), "GRO-099": ("01175201", "Per Kilogram"), "GRO-100": ("01172101", "Per Kilogram"),
 "GRO-101": ("01171401", "Each"), "GRO-102": ("01172201", "Each"), "GRO-104": ("01174901", "1 Kilogram"),
 "GRO-105": ("01174201", "Per Kilogram"), "GRO-107": ("01174501", "250 Gram"), "GRO-108": ("01178001", "1 Kilogram"),
 "GRO-111": ("01181101", "2.5 Kilogram"), "GRO-112": ("01181102", "2 Kilogram"), "GRO-113": ("01193101", "1 Kilogram"),
 "GRO-114": ("01151101", "2 Litre"), "GRO-116": ("01151301", "500 Millilitre"), "GRO-117": ("01179202", "410 Gram"),
 "GRO-121": ("01176101", "500 Gram"), "GRO-127": ("01194002", "100 Gram"), "GRO-129": ("01193903", "700 Millilitre"),
 "GRO-130": ("01193904", "750 Gram"), "GRO-131": ("01193902", "470 Gram"), "GRO-132": ("01169901", "400 Gram"),
 "GRO-133": ("01193901", "750 Millilitre"), "GRO-134": ("01184001", "400 Gram"), "GRO-135": ("01183901", "450 Gram"),
 "GRO-136": ("01183101", "375 Gram"), "GRO-138": ("01199002", "200 Gram"), "GRO-139": ("01199001", "10 Gram"),
 "GRO-140": ("01147002", "1 Litre"), "GRO-148": ("01210001", "1 Litre"), "GRO-150": ("01210002", "1 Litre"),
 "GRO-151": ("01260002", "2 Litre"), "GRO-153": ("01260001", "300 Millilitre"), "GRO-156": ("01147001", "1 Litre"),
 "GRO-157": ("01179901", "120 Gram"), "GRO-159": ("01185101", "80 Gram"), "GRO-160": ("01113901", "200 Gram"),
 "GRO-163": ("01189901", "100 Gram"), "GRO-164": ("01186001", "2 Litre"), "GRO-168": ("01192101", "900 Gram"),
 "GRO-170": ("01192301", "125 Millilitre"), "GRO-171": ("09322101", "8 Kilogram"), "GRO-173": ("09322102", "775 Gram"),
 "HHC-001": ("05611102", "2 Kilogram"), "HHC-004": ("05611106", "2 Litre"), "HHC-005": ("05611101", "500 Gram"),
 "HHC-006": ("05611103", "750 Millilitre"), "HHC-010": ("05611105", "500 Millilitre"), "HHC-014": ("13120204", "Pack Of 2 Ply 9 Rolls"),
 "HHC-016": ("05619201", "For 20"), "HHC-020": ("05619202", "Pack Of 6"), "PER-001": ("13120203", "175 Gram"),
 "PER-002": ("13120202", "400 Millilitre"), "PER-003": ("13120205", "100 Millilitre"), "PER-006": ("13120301", "150 Millilitre"),
 "PER-012": ("13120201", "400 Millilitre"), "PER-014": ("13120302", "400 Millilitre"), "HLT-001": ("06111201", "For 24"),
 "HLT-003": ("06111202", "100 Millilitre"), "HLT-006": ("06111203", "For 30"), "ALC-003": ("02122101", "6 X 330 Millilitre"),
 "ALC-005": ("02121102", "750 Millilitre"), "ALC-007": ("02110101", "750 Millilitre"), "ALC-008": ("02110102", "750 Millilitre"),
 "ALC-009": ("02110105", "750 Millilitre"), "ALC-010": ("02301101", "For 20"),
}
PROV = {"Western Cape": "WC", "Eastern Cape": "EC", "Northern Cape": "NC", "Free State": "FS", "KwaZulu-Natal": "KZN", "Kwa-Zulu Natal": "KZN",
        "North West": "NW", "Gauteng": "GP", "Mpumalanga": "MP", "Limpopo": "LP"}

def num(v):
    try: x = float(str(v).replace(",", ".")); return x if x > 0 else None
    except (TypeError, ValueError): return None

def month(col):  # 'M201701' -> '2017-01-15'
    return f"{col[1:5]}-{col[5:7]}-15"

au_path, pr_path, items_path = sys.argv[1:4]
rows = list(load_workbook(au_path, read_only=True).worksheets[0].iter_rows(values_only=True))
import re
MONTH_COL = re.compile(r"^M\d{6}$")  # skips extra columns such as 'March Online Price' (2020 lockdown)
hdr = rows[0]; mcols = [i for i, h in enumerate(hdr) if MONTH_COL.match(str(h))]
au = {(str(r[4]), str(r[9]).strip()): r for r in rows[1:]}
prow = list(load_workbook(pr_path, read_only=True).worksheets[0].iter_rows(values_only=True))
phdr = prow[0]; pm = [i for i, h in enumerate(phdr) if MONTH_COL.match(str(h))]
prov = {}
for r in prow[1:]:
    prov.setdefault((str(r[4]), str(r[9]).strip()), {})[PROV.get(str(r[10]).strip(), r[10])] = r

meta = {}
for r in list(load_workbook(items_path, read_only=True)["Items"].iter_rows(values_only=True))[1:]:
    if r[0] in MAP: meta[r[0]] = {"name": r[3], "cat": r[1], "sub": r[2], "spec": r[4], "unit": r[5], "tier": r[6]}

latest_col = hdr[mcols[-1]]
out = {"generated": datetime.date.today().isoformat(), "latestMonth": month(latest_col)[:7],
       "sources": ["Statistics South Africa, CPI average prices (all urban areas and provinces)"],
       "short": ["Stats SA average"], "series": {}, "items": {}, "prov": {}}
skipped = []
for iid, key in MAP.items():
    r = au.get(key)
    if r is None: skipped.append((iid, "no series")); continue
    pts = [[month(hdr[i]), round(v, 2), 0] for i in mcols if (v := num(r[i])) is not None]
    if not pts or pts[-1][0][:7] < month(latest_col)[:7]:
        skipped.append((iid, "series ended")); continue
    out["series"][iid] = pts
    out["items"][iid] = {**meta[iid], "statssa": f"{r[5]}, {key[1]}"}
    p = prov.get(key, {})
    cur = {c: num(pr[pm[-1]]) for c, pr in p.items()}
    ago = {c: num(pr[pm[-13]]) for c, pr in p.items()} if len(pm) >= 13 else {}
    if sum(1 for v in cur.values() if v) >= 5:
        out["prov"][iid] = {c: [cur[c], ago.get(c)] for c in sorted(cur) if cur[c]}
json.dump(out, open("public/data/statssa-history.json", "w"), separators=(",", ":"))
print(f"{len(out['series'])} items with history to {out['latestMonth']}, {len(out['prov'])} with provincial prices; skipped: {skipped}")
